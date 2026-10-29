"""Streaming adapters. Original records and stable logical locations are never discarded."""

import csv
import hashlib
import json
import zipfile
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import ijson
from openpyxl import load_workbook

from services.core.domain import money
from services.core.file_contract import KEYS, REQUIRED
from services.core.g1.common import canonical

VERSION = "normalized-v2.0"


def inspect_structure(path, entities):
    if path.suffix == ".json":
        seen = set()
        arrays = set()
        maps = []
        with path.open("rb") as f:
            depth = 0
            for prefix, event, value in ijson.parse(f):
                if event in {"start_map", "start_array"}:
                    depth += 1
                    if depth > 64:
                        raise ValueError("JSON_NESTING_LIMIT")
                elif event in {"end_map", "end_array"}:
                    depth -= 1
                if event == "start_map":
                    maps.append(set())
                elif event == "map_key":
                    if value in maps[-1]:
                        raise ValueError("DUPLICATE_JSON_KEY")
                    maps[-1].add(value)
                elif event == "end_map":
                    maps.pop()
                if prefix == "" and event == "map_key":
                    if value in seen:
                        raise ValueError("DUPLICATE_ENTITY_ARRAY")
                    seen.add(value)
                if prefix in entities and event == "start_array":
                    arrays.add(prefix)
        if seen != set(entities) or arrays != set(entities):
            raise ValueError("ENTITY_ARRAY_MISMATCH")
    elif path.suffix == ".xlsx":
        inspect_xlsx(path)
        workbook = load_workbook(path, read_only=True, keep_links=False)
        try:
            if set(workbook.sheetnames) != set(entities):
                raise ValueError("WORKSHEET_SET_MISMATCH")
        finally:
            workbook.close()


def records(path: Path, kind: str):
    if path.suffix == ".csv":
        with path.open(encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []
            if len(headers) != len(set(headers)) or set(headers) - set(REQUIRED[kind]) - {
                "operation"
            }:
                raise ValueError("INVALID_COLUMNS")
            for pos, row in enumerate(reader, 2):
                yield pos, row, {"file": path.name, "record": pos}
    elif path.suffix == ".json":
        with path.open("rb") as f:
            for pos, row in enumerate(ijson.items(f, f"{kind}.item"), 1):
                yield pos, row, {"file": path.name, "json_path": f"/{kind}/{pos - 1}"}
    elif path.suffix == ".xlsx":
        inspect_xlsx(path)
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if kind not in workbook.sheetnames:
                return
            sheet = workbook[kind]
            iterator = sheet.iter_rows()
            first = next(iterator, None)
            if first is None:
                return
            headers = [c.value for c in first]
            if (
                any(not isinstance(h, str) for h in headers)
                or len(headers) != len(set(headers))
                or set(headers) - set(REQUIRED[kind]) - {"operation"}
            ):
                raise ValueError("INVALID_COLUMNS")
            for pos, cells in enumerate(iterator, 2):
                if all(c.value is None for c in cells):
                    continue
                row = {k: c.value for k, c in zip(headers, cells, strict=True)}
                if any(c.data_type == "f" for c in cells):
                    row["_format_error"] = "FORMULA_FORBIDDEN"
                yield pos, row, {"file": path.name, "sheet": kind, "row": pos}
        finally:
            workbook.close()
    else:
        raise ValueError("UNSUPPORTED_FORMAT")


def inspect_xlsx(path):
    with zipfile.ZipFile(path) as z:
        entries = z.infolist()
        if len(entries) > 2000 or sum(e.file_size for e in entries) > 256 * 1024 * 1024:
            raise ValueError("XLSX_EXPANSION_LIMIT")
        for entry in entries:
            if entry.file_size > max(1, entry.compress_size) * 200 or any(
                s in entry.filename.lower() for s in ["vbaproject", "externallinks/"]
            ):
                raise ValueError("XLSX_UNSAFE_CONTENT")
        if any(
            ".." in e.filename.replace("\\", "/").split("/")
            or e.filename.replace("\\", "/").startswith("/")
            or ":" in e.filename.split("/")[0]
            for e in entries
        ):
            raise ValueError("XLSX_UNSAFE_CONTENT")


def normalize(kind, raw, fmt, mode):
    if not isinstance(raw, dict) or None in raw:
        raise ValueError("RECORD_NOT_OBJECT")
    if raw.get("_format_error"):
        raise ValueError(raw["_format_error"])
    if set(raw) - set(REQUIRED[kind]) - {"operation"}:
        raise ValueError("UNKNOWN_FIELD")
    operation = raw.get("operation", "upsert")
    if operation not in {"upsert", "delete"} or (operation == "delete" and mode != "delta"):
        raise ValueError("INVALID_OPERATION")
    if operation == "delete" and kind not in {"inventory", "rates"}:
        raise ValueError("MASTER_DELETE_NOT_SUPPORTED")
    data = {}
    needed = KEYS[kind] if operation == "delete" else REQUIRED[kind]
    for field in needed:
        value = raw.get(field)
        if value is None or value == "":
            raise ValueError("REQUIRED_VALUE:" + field)
        if field.endswith("_code") and not isinstance(value, str):
            raise ValueError("CODE_MUST_BE_TEXT:" + field)
        if field == "net_amount" and fmt == "json" and not isinstance(value, str):
            raise ValueError("JSON_AMOUNT_MUST_BE_TEXT")
        if field == "cancellation":
            policy = json.loads(value) if isinstance(value, str) else value
            if (
                not isinstance(policy, dict)
                or policy.get("kind") not in {"FLEX", "NONREF"}
                or policy.get("penalty_basis") != "stay_total"
                or type(policy.get("deadline_hours")) is not int
                or policy["deadline_hours"] < 0
            ):
                raise ValueError("INVALID_CANCELLATION")
            value = canonical(policy)
        elif isinstance(value, bool):
            value = "true" if value else "false"
        elif isinstance(value, (date, datetime)):
            value = value.isoformat()
        else:
            value = str(value)
        if field in {"quantity", "adults", "max_adults"}:
            if (
                not value.isascii()
                or not value.isdecimal()
                or int(value) < (0 if field == "quantity" else 1)
            ):
                raise ValueError("INVALID_INTEGER:" + field)
            value = str(int(value))
        if field == "adults" and value != "2":
            raise ValueError("UNSUPPORTED_OCCUPANCY")
        if field == "net_amount":
            value = format(money(value), ".4f")
        if field == "currency" and value != "USD":
            raise ValueError("UNSUPPORTED_CURRENCY")
        if field == "stay_date":
            date.fromisoformat(value)
        if field == "observed_at":
            parsed = datetime.fromisoformat(value)
            if parsed.utcoffset() is None:
                raise ValueError("TIMEZONE_REQUIRED")
        if field == "timezone":
            ZoneInfo(value)
        if field == "stop_sell" and value not in {"true", "false"}:
            raise ValueError("INVALID_BOOLEAN")
        if field == "bed_type" and value not in {"KING", "TWIN", "UNKNOWN"}:
            raise ValueError("INVALID_BED")
        if field == "meal" and value not in {"BREAKFAST_2", "NONE", "UNKNOWN"}:
            raise ValueError("INVALID_MEAL")
        if field == "payment" and value not in {"AT_PROPERTY", "PREPAID", "UNKNOWN"}:
            raise ValueError("INVALID_PAYMENT")
        if field == "tax_basis" and value not in {"NET_EXCLUSIVE", "GROSS_INCLUSIVE", "UNKNOWN"}:
            raise ValueError("INVALID_TAX_BASIS")
        data[field] = value
    data["operation"] = operation
    return data


def natural_key(kind, data):
    return canonical([data[k] for k in KEYS[kind]])


def partition(data):
    return data.get("stay_date", "MASTER") if isinstance(data, dict) else "*"


def file_hash(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        while block := f.read(1024 * 1024):
            h.update(block)
    return h.hexdigest()
