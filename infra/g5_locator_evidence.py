"""Real HTTP parsing and independent original-record checks for all three formats."""

import csv
import hashlib
import json
from datetime import UTC, datetime
from uuid import uuid4

from openpyxl import load_workbook

from infra.g5_http import Scenario
from infra.manage import ROOT, uid
from services.core.db import Database
from services.core.g1.common import digest
from services.core.g1.ingestion import STORE
from tests.g1_helpers import ok
from tests.g2_helpers import files


def main():
    s = Scenario(days=3)
    db = Database(s.settings.database_url)
    db.pool.open()
    report = {"at": datetime.now(UTC).isoformat(), "status": "IN_PROGRESS", "cases": []}
    try:
        for fmt in ["csv", "xlsx", "json"]:
            def edit(data, manifest):
                data["properties"][0]["name"] = "G5 original\nlogical record"

            batch = ok(s.client.post("/v2/imports", headers=s.headers(), files=files(fmt=fmt, code="G5-LOC-" + uuid4().hex[:6], edit=edit)))
            run = s.post("/v2/imports/" + batch["id"] + "/runs")
            ready = s.wait("/v2/processing-runs/" + run["id"], lambda x: x["status"] == "READY_TO_ACTIVATE")
            with db.transaction(uid("tenant-A")) as c:
                row = c.execute("SELECT r.position,r.locator,r.raw_hash,a.storage_key,a.sha256 FROM processing_record r JOIN artifact a ON a.id=r.artifact_id WHERE r.run_id=%s AND r.kind='properties' ORDER BY r.position LIMIT 1", (run["id"],)).fetchone()
            path = (STORE / row["storage_key"]).resolve()
            assert path.is_relative_to(STORE.resolve())
            raw_bytes = path.read_bytes()
            assert hashlib.sha256(raw_bytes).hexdigest() == row["sha256"]
            locator = row["locator"]
            if fmt == "csv":
                with path.open(encoding="utf-8", newline="") as stream:
                    records = list(csv.DictReader(stream))
                raw = records[locator["record"] - 2]
                assert len(raw_bytes.splitlines()) > len(records) + 1
                assert "\n" in raw["name"]
            elif fmt == "xlsx":
                book = load_workbook(path, read_only=True, data_only=False)
                try:
                    values = book[locator["sheet"]].iter_rows(values_only=True)
                    headers = next(values)
                    raw = dict(zip(headers, list(values)[locator["row"] - 2], strict=True))
                finally:
                    book.close()
            else:
                name, index = locator["json_path"].strip("/").split("/")
                raw = json.loads(raw_bytes)[name][int(index)]
            assert digest(raw) == row["raw_hash"]
            report["cases"].append({"format": fmt, "processing_id": ready["id"], "locator": locator, "artifact_sha256": row["sha256"], "record_hash_matches": True, "multiline_preserved": "\n" in raw["name"], "physical_csv_lines": len(raw_bytes.splitlines()) if fmt == "csv" else None})
        report["status"] = "PASSED"
    finally:
        db.pool.close()
        s.close()
        (ROOT / "docs/acceptance/g5-original-locators.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print("Three real HTTP format parses matched independently read original records and hashes")


if __name__ == "__main__":
    main()
